import pandas as pd
import openpyxl

# Load the Excel file
file_path = r"C:\Users\kit02\OneDrive\研究\エクセル\x-means_\spot_summary_2024_ABO_O_normalized_mean48_xmeans_daytrend_seed42_kmax50_spot13_yaxis_fix_20260317.xlsx"

# Load the workbook
wb = openpyxl.load_workbook(file_path)

# List all sheet names
print("Sheet names:", wb.sheetnames)

# Check the dates sheet
dates_sheet = "spot_summary_2024 (1)_dates"
if dates_sheet in wb.sheetnames:
    df_dates = pd.read_excel(file_path, sheet_name=dates_sheet)
    print(f"\nSheet '{dates_sheet}' loaded.")
    print("DataFrame shape:", df_dates.shape)
    print("Columns:", df_dates.columns.tolist())
    print("First 5 rows:")
    print(df_dates.head())
    
    # Check for cluster_id and date columns
    if 'cluster_id' in df_dates.columns and 'date' in df_dates.columns:
        print("\nCluster ID to Date mapping:")
        for _, row in df_dates.iterrows():
            print(f"Cluster {row['cluster_id']}: {row['date']}")
    else:
        print("\nNo 'cluster_id' or 'date' columns found in dates sheet.")
else:
    print(f"Sheet '{dates_sheet}' not found.")

# Focus on the specific sheet
sheet_name = "spot_summary_2024 (1)_spot"
if sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\nSheet '{sheet_name}' loaded.")
    
    # Read the data into a DataFrame
    df = pd.read_excel(file_path, sheet_name=sheet_name)
    print("DataFrame shape:", df.shape)
    print("Columns:", df.columns.tolist())
    print("First 5 rows:")
    print(df.head())
    
    # Check for cluster_id column
    if 'cluster_id' in df.columns:
        print("\nCluster IDs:")
        print(df['cluster_id'].unique())
    else:
        print("\nNo 'cluster_id' column found.")
        
    # Check for date column
    date_cols = [col for col in df.columns if 'date' in col.lower()]
    if date_cols:
        print(f"\nDate columns: {date_cols}")
        for col in date_cols:
            print(f"Unique values in {col}: {df[col].unique()[:10]}")  # First 10 unique values
    else:
        print("\nNo date columns found.")
        
else:
    print(f"Sheet '{sheet_name}' not found.")
import argparse
import csv
import math
import os
import random
import re
from typing import List, Optional, Sequence, Tuple

import numpy as np
import openpyxl


DEFAULT_INPUT_NAME = "extracted_A_L_yyyymm_split_1-0_2-1_D_normalized_minmax.xlsx"


def find_file_under_onedrive(filename: str) -> Optional[str]:
    start = r"C:\Users\kit02\OneDrive"
    for root, _, files in os.walk(start):
        if filename in files:
            return os.path.join(root, filename)
    return None


def parse_date(c_value: object) -> Optional[str]:
    if c_value is None:
        return None
    text = str(c_value).strip()
    m = re.match(r"^(\d{8})", text)
    return m.group(1) if m else None


def format_date_yyyy_mm_dd(s: str) -> str:
    if len(s) == 8 and s.isdigit():
        return f"{s[0:4]}-{s[4:6]}-{s[6:8]}"
    return s


def parse_float(v: object) -> Optional[float]:
    if v is None:
        return None
    text = str(v).strip()
    if text == "":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def kmeans_pp_init(data: np.ndarray, k: int, rng: random.Random) -> np.ndarray:
    n = data.shape[0]
    first = rng.randrange(n)
    centers = [data[first]]

    while len(centers) < k:
        d2 = np.min(
            np.stack([np.sum((data - c) ** 2, axis=1) for c in centers], axis=1),
            axis=1,
        )
        total = float(np.sum(d2))
        if total <= 0.0:
            centers.append(data[rng.randrange(n)])
            continue
        r = rng.random() * total
        acc = 0.0
        idx = 0
        for i, val in enumerate(d2):
            acc += float(val)
            if acc >= r:
                idx = i
                break
        centers.append(data[idx])

    return np.asarray(centers, dtype=float)


def kmeans_nd(data: np.ndarray, k: int, rng: random.Random, max_iter: int = 100) -> Tuple[np.ndarray, np.ndarray]:
    n, d = data.shape
    if n == 0:
        return np.empty((0,), dtype=int), np.empty((0, d), dtype=float)
    if k <= 1:
        center = np.mean(data, axis=0, keepdims=True)
        return np.zeros(n, dtype=int), center

    centers = kmeans_pp_init(data, k, rng)
    labels = np.zeros(n, dtype=int)

    for _ in range(max_iter):
        dist = np.stack([np.sum((data - c) ** 2, axis=1) for c in centers], axis=1)
        new_labels = np.argmin(dist, axis=1)
        changed = np.any(new_labels != labels)
        labels = new_labels

        new_centers = centers.copy()
        for c in range(k):
            idx = np.where(labels == c)[0]
            if idx.size > 0:
                new_centers[c] = np.mean(data[idx], axis=0)
            else:
                new_centers[c] = data[rng.randrange(n)]

        if np.allclose(new_centers, centers, atol=1e-12) and not changed:
            centers = new_centers
            break
        centers = new_centers

    # Stable ordering by centroid lexicographic order.
    order = np.lexsort(np.flipud(centers.T))
    remap = {int(old): int(new) for new, old in enumerate(order)}
    centers = centers[order]
    labels = np.asarray([remap[int(x)] for x in labels], dtype=int)
    return labels, centers


def bic_score_nd(data: np.ndarray, labels: np.ndarray, centers: np.ndarray) -> float:
    n, d = data.shape
    k = centers.shape[0]
    if n == 0 or k == 0:
        return float("-inf")

    sse = 0.0
    counts = np.zeros(k, dtype=int)
    for i in range(n):
        c = int(labels[i])
        counts[c] += 1
        diff = data[i] - centers[c]
        sse += float(np.dot(diff, diff))

    dof = max(n - k, 1)
    var = max(sse / dof, 1e-12)
    c0 = -0.5 * d * math.log(2.0 * math.pi * var)

    loglik = 0.0
    for i in range(n):
        c = int(labels[i])
        pi_k = max(counts[c] / n, 1e-12)
        diff = data[i] - centers[c]
        dist2 = float(np.dot(diff, diff))
        loglik += c0 - dist2 / (2.0 * var) + math.log(pi_k)

    p = k * d + (k - 1) + 1
    return loglik - 0.5 * p * math.log(n)


def xmeans_nd(data: np.ndarray, rng: random.Random, kmax: int = 8) -> Tuple[np.ndarray, np.ndarray]:
    n, d = data.shape
    if n == 0:
        return np.empty((0,), dtype=int), np.empty((0, d), dtype=float)
    if n == 1:
        return np.zeros(1, dtype=int), data.copy()

    labels, centers = kmeans_nd(data, 1, rng)

    while True:
        k = centers.shape[0]
        if k >= kmax:
            break

        split_any = False
        next_centers: List[np.ndarray] = []

        for c in range(k):
            idx = np.where(labels == c)[0]
            cluster_data = data[idx]

            if cluster_data.shape[0] < 4:
                next_centers.append(centers[c])
                continue

            parent_labels = np.zeros(cluster_data.shape[0], dtype=int)
            parent_centers = np.mean(cluster_data, axis=0, keepdims=True)
            parent_bic = bic_score_nd(cluster_data, parent_labels, parent_centers)

            child_labels, child_centers = kmeans_nd(cluster_data, 2, rng)
            child_bic = bic_score_nd(cluster_data, child_labels, child_centers)

            if child_bic > parent_bic and (len(next_centers) + (k - c) + 1) <= kmax:
                next_centers.extend([child_centers[0], child_centers[1]])
                split_any = True
            else:
                next_centers.append(centers[c])

        if not split_any:
            break

        _, centers = kmeans_nd(np.asarray(data), len(next_centers), rng)
        labels, centers = kmeans_nd(np.asarray(data), len(next_centers), rng)

    return labels, centers


def serialize_vec(v: np.ndarray) -> str:
    return "[" + ",".join(f"{float(x):.8f}" for x in v.tolist()) + "]"


def run(input_path: str, assign_csv: str, rep_csv: str, seed: int, kmax: int) -> Tuple[int, int]:
    rng = random.Random(seed)
    wb = openpyxl.load_workbook(input_path, data_only=True, read_only=True)

    assign_rows = [["sheet", "date", "cluster_id", "distance_to_centroid", "num_points", "mean_d"]]
    rep_rows = [["sheet", "cluster_id", "centroid_vector", "nearest_date", "distance", "members_count"]]

    total_dates = 0
    total_clusters = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        groups = {}
        for row in ws.iter_rows(values_only=True):
            date = parse_date(row[2] if len(row) >= 3 else None)
            if date is None:
                continue
            dval = parse_float(row[3] if len(row) >= 4 else None)
            if dval is None:
                continue
            groups.setdefault(date, []).append(dval)

        if not groups:
            continue

        dates = sorted(groups.keys())
        max_len = max(len(groups[d]) for d in dates)
        if max_len == 0:
            continue

        day_vectors = []
        means = []
        for d in dates:
            vals = sorted(groups[d])
            means.append(float(np.mean(vals)))
            if len(vals) < max_len:
                vals = vals + [vals[-1]] * (max_len - len(vals))
            day_vectors.append(vals)
        data = np.asarray(day_vectors, dtype=float)

        labels, centers = xmeans_nd(data, rng=rng, kmax=kmax)
        total_dates += data.shape[0]
        total_clusters += centers.shape[0]

        for i, date in enumerate(dates):
            c = int(labels[i])
            dist = float(np.linalg.norm(data[i] - centers[c]))
            assign_rows.append(
                [sheet_name, format_date_yyyy_mm_dd(date), c, dist, len(groups[date]), means[i]]
            )

        for c in range(centers.shape[0]):
            idx = np.where(labels == c)[0]
            if idx.size == 0:
                continue
            dists = np.linalg.norm(data[idx] - centers[c], axis=1)
            best_local = int(np.argmin(dists))
            best_idx = int(idx[best_local])
            rep_rows.append(
                [
                    sheet_name,
                    c,
                    serialize_vec(centers[c]),
                    format_date_yyyy_mm_dd(dates[best_idx]),
                    float(dists[best_local]),
                    int(idx.size),
                ]
            )

    with open(assign_csv, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(assign_rows)
    with open(rep_csv, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rep_rows)

    return total_dates, total_clusters


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Cluster by day-profile (all D values in a day) and find nearest date to each centroid."
        )
    )
    parser.add_argument("--input", dest="input_path", default=None, help="Input normalized xlsx path")
    parser.add_argument("--assign", dest="assign_csv", default=None, help="Output per-date assignment csv path")
    parser.add_argument("--rep", dest="rep_csv", default=None, help="Output per-cluster representative-date csv path")
    parser.add_argument("--seed", type=int, default=42, help="Random seed")
    parser.add_argument("--kmax", type=int, default=8, help="Maximum clusters")
    args = parser.parse_args()

    input_path = args.input_path
    if input_path is None:
        found = find_file_under_onedrive(DEFAULT_INPUT_NAME)
        if found is None:
            raise SystemExit(f"Input not found under OneDrive: {DEFAULT_INPUT_NAME}")
        input_path = found

    base = os.path.dirname(input_path)
    assign_csv = args.assign_csv or os.path.join(
        base, f"day_profile_xmeans_seed{args.seed}_assignments.csv"
    )
    rep_csv = args.rep_csv or os.path.join(
        base, f"day_profile_xmeans_seed{args.seed}_representatives.csv"
    )

    total_dates, total_clusters = run(
        input_path=input_path,
        assign_csv=assign_csv,
        rep_csv=rep_csv,
        seed=args.seed,
        kmax=args.kmax,
    )
    print(f"input={input_path}")
    print(f"assignments={assign_csv}")
    print(f"representatives={rep_csv}")
    print(f"seed={args.seed}")
    print(f"kmax={args.kmax}")
    print(f"total_dates={total_dates}")
    print(f"total_clusters={total_clusters}")


if __name__ == "__main__":
    main()

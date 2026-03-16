from __future__ import annotations

import argparse
import csv
import shutil
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl


DEFAULT_CLUSTER_RESULT_NAME = "spot_summary_2024_ABO_O_normalized_mean48_xmeans_daytrend_seed42_kmax50.xlsx"
DEFAULT_RAW_NAME = "spot_summary_2024 (1).xlsx"


def find_under_onedrive(filename: str) -> Optional[Path]:
    start = Path(r"C:/Users/kit02/OneDrive")
    for p in start.rglob(filename):
        if p.is_file():
            return p
    return None


def parse_date_key(v: object) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if s == "":
        return None
    digits = "".join(ch for ch in s if ch.isdigit())
    if len(digits) >= 8:
        return f"{digits[0:4]}-{digits[4:6]}-{digits[6:8]}"
    return s


def parse_float(v: object) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_int(v: object) -> Optional[int]:
    f = parse_float(v)
    if f is None:
        return None
    return int(f)


def copy_to_local(path: Path, local_name: str) -> Path:
    local = Path.home() / local_name
    if path.resolve() == local.resolve():
        return local
    shutil.copy2(path, local)
    return local


def load_cluster_info(cluster_wb: openpyxl.Workbook) -> Tuple[Dict[int, int], Dict[int, Tuple[str, float]], int, str]:
    assign_ws = None
    repr_ws = None
    source_sheet = None
    for ws in cluster_wb.worksheets:
        if ws.title.endswith("_assign"):
            assign_ws = ws
        elif ws.title.endswith("_repr"):
            repr_ws = ws

    if assign_ws is None or repr_ws is None:
        raise RuntimeError("assign/repr sheets were not found in cluster result workbook.")

    member_counts: Dict[int, int] = defaultdict(int)
    for row in assign_ws.iter_rows(min_row=2, values_only=True):
        sheet_name = row[0] if len(row) > 0 else None
        cluster_id = parse_int(row[2] if len(row) > 2 else None)
        if source_sheet is None and sheet_name is not None:
            source_sheet = str(sheet_name)
        if cluster_id is None:
            continue
        member_counts[cluster_id] += 1

    repr_map: Dict[int, Tuple[str, float]] = {}
    for row in repr_ws.iter_rows(min_row=2, values_only=True):
        cluster_id = parse_int(row[1] if len(row) > 1 else None)
        rep_date = parse_date_key(row[2] if len(row) > 2 else None)
        dist = parse_float(row[4] if len(row) > 4 else None)
        if cluster_id is None or rep_date is None or dist is None:
            continue
        repr_map[cluster_id] = (rep_date, dist)

    total_clusters = len(member_counts)
    if source_sheet is None:
        source_sheet = ""
    return member_counts, repr_map, total_clusters, source_sheet


def load_raw_series_by_date(raw_wb: openpyxl.Workbook, source_sheet: str) -> Dict[str, List[Optional[float]]]:
    # Prefer exact sheet; fallback to first sheet.
    if source_sheet and source_sheet in raw_wb.sheetnames:
        ws = raw_wb[source_sheet]
    else:
        ws = raw_wb[raw_wb.sheetnames[0]]

    tmp: Dict[str, List[Tuple[int, Optional[float]]]] = defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = parse_date_key(row[0] if len(row) > 0 else None)  # A
        tc = parse_int(row[1] if len(row) > 1 else None)      # B
        val = parse_float(row[14] if len(row) > 14 else None) # O
        if d is None or tc is None:
            continue
        tmp[d].append((tc, val))

    out: Dict[str, List[Optional[float]]] = {}
    for d, points in tmp.items():
        points_sorted = sorted(points, key=lambda x: x[0])
        series = [v for _tc, v in points_sorted]
        # Keep first 48 points; if shorter, pad with None.
        series = series[:48]
        if len(series) < 48:
            series = series + [None] * (48 - len(series))
        out[d] = series
    return out


def run(cluster_path: Path, raw_path: Path, out_xlsx: Path, out_csv: Path) -> None:
    local_cluster = copy_to_local(cluster_path, "cluster_result_for_summary.xlsx")
    local_raw = copy_to_local(raw_path, "raw_spot_summary_for_summary.xlsx")

    cluster_wb = openpyxl.load_workbook(local_cluster, data_only=True, read_only=True)
    raw_wb = openpyxl.load_workbook(local_raw, data_only=True, read_only=True)

    member_counts, repr_map, total_clusters, source_sheet = load_cluster_info(cluster_wb)
    raw_series_by_date = load_raw_series_by_date(raw_wb, source_sheet)

    headers = [
        "total_clusters",
        "cluster_id",
        "member_days",
        "representative_date",
        "distance_to_center",
    ] + [f"raw_t{i}" for i in range(1, 49)]

    rows: List[List[object]] = []
    for c in sorted(member_counts.keys()):
        member_days = member_counts[c]
        rep_date, dist = repr_map.get(c, (None, None))
        raw_series = raw_series_by_date.get(rep_date, [None] * 48) if rep_date else [None] * 48
        rows.append([total_clusters, c, member_days, rep_date, dist] + raw_series)

    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = "cluster_summary"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb_out.save(out_xlsx)

    with out_csv.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="Summarize X-means day-trend clustering results.")
    parser.add_argument("--cluster", default=None, help="Cluster result xlsx path")
    parser.add_argument("--raw", default=None, help="Raw pre-normalization xlsx path")
    parser.add_argument("--out-xlsx", default=None, help="Output summary xlsx path")
    parser.add_argument("--out-csv", default=None, help="Output summary csv path")
    args = parser.parse_args()

    cluster_path = Path(args.cluster) if args.cluster else find_under_onedrive(DEFAULT_CLUSTER_RESULT_NAME)
    raw_path = Path(args.raw) if args.raw else find_under_onedrive(DEFAULT_RAW_NAME)

    if cluster_path is None:
        raise FileNotFoundError(f"Cluster workbook not found: {DEFAULT_CLUSTER_RESULT_NAME}")
    if raw_path is None:
        raise FileNotFoundError(f"Raw workbook not found: {DEFAULT_RAW_NAME}")

    base_dir = cluster_path.parent
    out_xlsx = Path(args.out_xlsx) if args.out_xlsx else base_dir / "spot_summary_2024_xmeans_daytrend_summary.xlsx"
    out_csv = Path(args.out_csv) if args.out_csv else base_dir / "spot_summary_2024_xmeans_daytrend_summary.csv"

    run(cluster_path=cluster_path, raw_path=raw_path, out_xlsx=out_xlsx, out_csv=out_csv)
    print(f"cluster={cluster_path}")
    print(f"raw={raw_path}")
    print(f"out_xlsx={out_xlsx}")
    print(f"out_csv={out_csv}")


if __name__ == "__main__":
    main()

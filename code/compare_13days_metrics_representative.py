from __future__ import annotations

import argparse
import math
import re
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import openpyxl


DEFAULT_DEG_PATH = r"C:\Users\kit02\OneDrive\研究\result\2d_piecewise_batch_2_20260306_161151\c01_13days_extracted_metrics.xlsx"
DEFAULT_NODEG_PATH = r"C:\Users\kit02\OneDrive\研究\result\2d_piecewise_batch_2_20260302_195041_nodeg\c01_13days_extracted_metrics_nodeg.xlsx"


def safe_sheet_name(name: str, used: set[str]) -> str:
    sanitized = re.sub(r"[\\/?*\\[\\]:]", "_", str(name))
    base = sanitized[:30]
    if base == "":
        base = "sheet"
    if base not in used:
        used.add(base)
        return base
    i = 2
    while True:
        suffix = f"_{i}"
        candidate = f"{base[:30-len(suffix)]}{suffix}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        i += 1


def normalize_date(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (int, float)) and math.isnan(v):
        return ""
    if hasattr(v, "strftime"):
        try:
            return v.strftime("%Y-%m-%d")
        except Exception:
            return str(v)
    text = str(v).strip()
    if text == "":
        return ""
    digits = "".join(ch for ch in text if ch.isdigit())
    if len(digits) >= 8:
        return f"{digits[0:4]}-{digits[4:6]}-{digits[6:8]}"
    return text


def _copy_for_read(path: Path) -> Tuple[Path, bool]:
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        wb.close()
        return path, False
    except PermissionError:
        tmp = Path(tempfile.gettempdir()) / f"comp_{path.stem}_{id(path)}.xlsx"
        ps = (
            f"Copy-Item -LiteralPath '{path}' -Destination '{tmp}' -Force; "
            f"if (-not (Test-Path -LiteralPath '{tmp}')) {{ throw 'copy failed' }}"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
        return tmp, True


def load_wb(path: Path) -> Tuple[openpyxl.Workbook, Path, bool]:
    copied_path, used_copy = _copy_for_read(path)
    wb = openpyxl.load_workbook(copied_path, data_only=True, read_only=True)
    return wb, copied_path, used_copy and copied_path != path


def parse_soc_rows(ws: openpyxl.worksheet.worksheet.Worksheet) -> Dict[str, Dict[str, float]]:
    rows = {}
    header = [h for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    t_cols = [i for i, h in enumerate(header) if isinstance(h, str) and re.fullmatch(r"t\d+", h)]

    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r:
            continue
        date = normalize_date(r[0])
        if date == "":
            continue
        values = {}
        for i in t_cols:
            key = header[i]
            values[key] = r[i] if (len(r) > i and r[i] is not None) else None
        rows[date] = values
    return rows


def parse_spot_rows(
    ws: openpyxl.worksheet.worksheet.Worksheet,
) -> Dict[str, Dict[str, Tuple[float, float]]]:
    rows = {}
    header = [h for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    buy_cols: List[Tuple[int, str]] = []
    sell_cols: List[Tuple[int, str]] = []
    for i, h in enumerate(header):
        if not isinstance(h, str):
            continue
        m = re.match(r"buy_t(\d+)", h)
        if m:
            buy_cols.append((i, m.group(1)))
            continue
        m = re.match(r"sell_t(\d+)", h)
        if m:
            sell_cols.append((i, m.group(1)))

    buy_map = {k: i for i, k in buy_cols}
    sell_map = {k: i for i, k in sell_cols}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r:
            continue
        date = normalize_date(r[0])
        if date == "":
            continue
        values: Dict[str, Tuple[float, float]] = {}
        for step in sorted(set(buy_map) & set(sell_map), key=lambda x: int(x)):
            b = r[buy_map[step]] if len(r) > buy_map[step] else None
            s = r[sell_map[step]] if len(r) > sell_map[step] else None
            values[step] = (b, s)
        rows[date] = values
    return rows


def parse_primary_secondary_rows(
    ws: openpyxl.worksheet.worksheet.Worksheet,
) -> Dict[str, Dict[str, Tuple[float, float]]]:
    rows = {}
    header = [h for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    p_cols: List[Tuple[int, str]] = []
    s_cols: List[Tuple[int, str]] = []
    for i, h in enumerate(header):
        if not isinstance(h, str):
            continue
        m = re.match(r"primary_b(\d+)", h)
        if m:
            p_cols.append((i, m.group(1)))
            continue
        m = re.match(r"secondary_b(\d+)", h)
        if m:
            s_cols.append((i, m.group(1)))

    p_map = {k: i for i, k in p_cols}
    s_map = {k: i for i, k in s_cols}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r:
            continue
        date = normalize_date(r[0])
        if date == "":
            continue
        values: Dict[str, Tuple[float, float]] = {}
        for step in sorted(set(p_map) & set(s_map), key=lambda x: int(x)):
            p = r[p_map[step]] if len(r) > p_map[step] else None
            s = r[s_map[step]] if len(r) > s_map[step] else None
            values[step] = (p, s)
        rows[date] = values
    return rows


def append_section_header(ws: openpyxl.worksheet.worksheet.Worksheet, title: str) -> None:
    ws.append([])
    ws.append([title])


def append_soc_section(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    deg: Dict[str, float],
    nodeg: Dict[str, float],
) -> None:
    append_section_header(ws, "SOC推移")
    ws.append(["step", "SOC_劣化あり", "SOC_劣化なし", "差分(あり-なし)"])
    keys = sorted(set(deg) | set(nodeg), key=lambda s: int(s[1:]))
    for k in keys:
        a = deg.get(k)
        b = nodeg.get(k)
        if a is None and b is None:
            continue
        diff = None if a is None or b is None else (a - b)
        ws.append([k, a, b, diff])


def append_spot_section(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    deg: Dict[str, Tuple[float, float]],
    nodeg: Dict[str, Tuple[float, float]],
) -> None:
    append_section_header(ws, "spot売買量")
    ws.append(
        [
            "step",
            "buy_劣化あり",
            "buy_劣化なし",
            "buy差分(あり-なし)",
            "sell_劣化あり",
            "sell_劣化なし",
            "sell差分(あり-なし)",
        ]
    )
    keys = sorted(set(deg) | set(nodeg), key=lambda s: int(s))
    for k in keys:
        db = deg.get(k)
        dn = nodeg.get(k)
        dbuy = db[0] if db is not None else None
        dsell = db[1] if db is not None else None
        nbuy = dn[0] if dn is not None else None
        nsell = dn[1] if dn is not None else None
        diff_buy = None if dbuy is None or nbuy is None else (dbuy - nbuy)
        diff_sell = None if dsell is None or nsell is None else (dsell - nsell)
        ws.append([k, dbuy, nbuy, diff_buy, dsell, nsell, diff_sell])


def append_ps_section(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    deg: Dict[str, Tuple[float, float]],
    nodeg: Dict[str, Tuple[float, float]],
) -> None:
    append_section_header(ws, "一次・二次売電量")
    ws.append(
        [
            "step",
            "1次_劣化あり",
            "1次_劣化なし",
            "1次差分(あり-なし)",
            "2次_劣化あり",
            "2次_劣化なし",
            "2次差分(あり-なし)",
        ]
    )
    keys = sorted(set(deg) | set(nodeg), key=lambda s: int(s))
    for k in keys:
        dp = deg.get(k)
        dn = nodeg.get(k)
        d1 = dp[0] if dp is not None else None
        d2 = dp[1] if dp is not None else None
        n1 = dn[0] if dn is not None else None
        n2 = dn[1] if dn is not None else None
        diff1 = None if d1 is None or n1 is None else (d1 - n1)
        diff2 = None if d2 is None or n2 is None else (d2 - n2)
        ws.append([k, d1, n1, diff1, d2, n2, diff2])


def read_metrics(path: Path) -> Dict[str, Dict[str, Dict[str, Any]]]:
    wb, copied_path, is_temp = load_wb(path)
    try:
        try:
            ws_soc = wb["soc_transition_wide"]
        except KeyError:
            raise KeyError(f"sheet not found: soc_transition_wide in {path}")
        try:
            ws_spot = wb["spot_buy_sell_wide"]
        except KeyError:
            raise KeyError(f"sheet not found: spot_buy_sell_wide in {path}")
        try:
            ws_ps = wb["primary_secondary_wide"]
        except KeyError:
            raise KeyError(f"sheet not found: primary_secondary_wide in {path}")

        return {
            "soc": parse_soc_rows(ws_soc),
            "spot": parse_spot_rows(ws_spot),
            "ps": parse_primary_secondary_rows(ws_ps),
        }
    finally:
        wb.close()
        if is_temp:
            try:
                copied_path.unlink()
            except Exception:
                pass


def normalize_metrics_data(
    path_degradation: Path,
    path_no_degradation: Path,
) -> Tuple[Dict[str, Any], set[str]]:
    deg = read_metrics(path_degradation)
    nodeg = read_metrics(path_no_degradation)

    dates = set(deg["soc"].keys()) & set(nodeg["soc"].keys())
    # Include all matching dates only.
    return {"deg": deg, "nodeg": nodeg}, dates


def build_summary_workbook(
    output: Path,
    metrics_deg: Dict[str, Dict[str, Dict[str, Any]]],
    metrics_nodeg: Dict[str, Dict[str, Dict[str, Any]]],
    dates: Iterable[str],
) -> None:
    wb_out = openpyxl.Workbook()
    ws_index = wb_out.active
    ws_index.title = "index"
    ws_index.append(["representative_date", "exists_deg", "exists_nodeg", "status"])

    used = {ws_index.title}
    for d in sorted(dates):
        # Existence check
        exists_deg = "YES"
        exists_nodeg = "YES"
        if d not in metrics_deg["soc"] or d not in metrics_nodeg["soc"]:
            exists_deg = "YES" if d in metrics_deg["soc"] else "NO"
            exists_nodeg = "YES" if d in metrics_nodeg["soc"] else "NO"
            ws_index.append([d, exists_deg, exists_nodeg, "skip"])
            continue

        ws_index.append([d, exists_deg, exists_nodeg, "ok"])

        ws = wb_out.create_sheet(safe_sheet_name(d, used))
        ws.append([f"代表日: {d}  (劣化あり vs 劣化なし)"])
        ws.append([])
        ws.append(["区分", "対象", "sheet"])

        append_soc_section(
            ws,
            {k: metrics_deg["soc"][d].get(k) for k in metrics_deg["soc"][d]},
            {k: metrics_nodeg["soc"][d].get(k) for k in metrics_nodeg["soc"][d]},
        )
        append_spot_section(
            ws,
            {k: metrics_deg["spot"][d].get(k) for k in metrics_deg["spot"][d]},
            {k: metrics_nodeg["spot"][d].get(k) for k in metrics_nodeg["spot"][d]},
        )
        append_ps_section(
            ws,
            {k: metrics_deg["ps"][d].get(k) for k in metrics_deg["ps"][d]},
            {k: metrics_nodeg["ps"][d].get(k) for k in metrics_nodeg["ps"][d]},
        )

    wb_out.save(output)


def run(deg_path: Path, nodeg_path: Path, output: Path) -> Path:
    data, dates = normalize_metrics_data(deg_path, nodeg_path)
    build_summary_workbook(output, data["deg"], data["nodeg"], dates)
    return output


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Compare degraded vs no-degradation extracted 13-day metrics per representative date."
    )
    p.add_argument("--deg", default=DEFAULT_DEG_PATH, help="Workbook with degradation")
    p.add_argument(
        "--nodeg",
        default=DEFAULT_NODEG_PATH,
        help="Workbook without degradation",
    )
    p.add_argument(
        "--out",
        default=None,
        help="Output XLSX path. Default: folder of deg file with suffix _deg_vs_nodeg_compare.xlsx",
    )
    return p.parse_args()


def main() -> None:
    args = parse_args()
    deg = Path(args.deg)
    nodeg = Path(args.nodeg)
    if not deg.exists():
        raise FileNotFoundError(f"deg file not found: {deg}")
    if not nodeg.exists():
        raise FileNotFoundError(f"nodeg file not found: {nodeg}")

    if args.out:
        output = Path(args.out)
    else:
        output = deg.parent / f"{deg.stem}_deg_vs_nodeg_compare.xlsx"

    out = run(deg, nodeg, output)
    print(f"output={out}")


if __name__ == "__main__":
    main()

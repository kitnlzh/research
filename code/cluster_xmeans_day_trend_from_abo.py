from __future__ import annotations

import argparse
import math
import os
import numpy as np
import random
import shutil
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple
import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.utils.cell import get_column_letter

from cluster_xmeans_day_profile import bic_score_nd, kmeans_nd

DEFAULT_INPUT_NAME = "spot_summary_2024_ABO_O_normalized_mean48.xlsx"


def find_file_under_onedrive(filename: str) -> Optional[Path]:
    start = Path(r"C:/Users/kit02/OneDrive")
    for p in start.rglob(filename):
        if p.is_file():
            return p
    return None


def parse_date_key(v: object) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    text = str(v).strip()
    if text == "":
        return None
    # Support YYYYMMDD and YYYY-MM-DD prefixes.
    digits = "".join(ch for ch in text if ch.isdigit())
    if len(digits) >= 8:
        y, m, d = digits[0:4], digits[4:6], digits[6:8]
        return f"{y}-{m}-{d}"
    return text


def parse_float(v: object) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return float(v)
    s = str(v).strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_int(v: object) -> Optional[int]:
    f = parse_float(v)
    if f is None:
        return None
    return int(f)


def add_cluster_time_series_charts(
    wb_out: openpyxl.Workbook,
    src_sheet: openpyxl.worksheet.worksheet.Worksheet,
    ws_center: openpyxl.worksheet.worksheet.Worksheet,
    n_clusters: int,
    max_len: int,
    labels: np.ndarray,
    day_dates: List[str],
    day_data: np.ndarray,
    used_sheet_names: set[str],
    chart_count: int,
    y_label: str,
    chart_suffix: str,
    charts_per_row: int = 3,
) -> None:
    if n_clusters <= 0 or max_len <= 0:
        return

    chart_count = min(n_clusters, chart_count)
    if chart_count <= 0:
        return

    chart_sheet_name = safe_sheet_name(f"{src_sheet.title}_{chart_suffix}", used_sheet_names)
    ws_chart = wb_out.create_sheet(chart_sheet_name)
    ws_chart.append(["cluster_id", "date", f"{y_label.lower()}_chart"])

    data_sheet_name = safe_sheet_name(f"{src_sheet.title}_{chart_suffix}_data", used_sheet_names)
    ws_data = wb_out.create_sheet(data_sheet_name)
    ws_data.append(["cluster_id", "date"] + [f"t{i+1}" for i in range(max_len)])
    ws_data.sheet_state = "hidden"

    ranges: Dict[int, Tuple[int, int]] = {}
    for c in range(chart_count):
        idxs = np.where(labels == c)[0]
        if idxs.size == 0:
            continue
        start_row = ws_data.max_row + 1
        for i in idxs:
            ws_data.append([c, day_dates[i]] + [float(v) for v in day_data[i].tolist()])
        end_row = ws_data.max_row
        ranges[c] = (start_row, end_row)

    x_categories = Reference(ws_center, min_col=3, min_row=1, max_col=2 + max_len, max_row=1)

    plotted = 0
    for c in range(chart_count):
        if c not in ranges:
            continue
        start_row, end_row = ranges[c]
        if start_row > end_row:
            continue

        values = Reference(
            ws_data,
            min_col=3,
            min_row=start_row,
            max_col=2 + max_len,
            max_row=end_row,
        )

        chart = LineChart()
        chart.title = f"{src_sheet.title} {y_label} Cluster {c}"
        chart.y_axis.title = y_label
        chart.height = 6.0
        chart.width = 11.0
        chart.legend = None
        chart.x_axis.delete = True
        chart.x_axis.title = None
        chart.x_axis.majorGridlines = None
        chart.x_axis.minorGridlines = None
        chart.x_axis.majorTickMark = "none"
        chart.x_axis.minorTickMark = "none"
        chart.x_axis.txPr = None
        chart.add_data(values, titles_from_data=False, from_rows=True)
        chart.set_categories(x_categories)

        ws_chart.append([c, f"members={end_row - start_row + 1}", f"plot {c}"])

        col = (plotted % charts_per_row) * 8 + 1
        row = (plotted // charts_per_row) * 18 + 2
        ws_chart.add_chart(chart, anchor=f"{get_column_letter(col)}{row}")
        plotted += 1


def add_soc_charts(
    wb_out: openpyxl.Workbook,
    src_sheet: openpyxl.worksheet.worksheet.Worksheet,
    ws_center: openpyxl.worksheet.worksheet.Worksheet,
    n_clusters: int,
    max_len: int,
    labels: np.ndarray,
    day_dates: List[str],
    day_data: np.ndarray,
    used_sheet_names: set[str],
    max_charts: int = 13,
    charts_per_row: int = 3,
) -> None:
    add_cluster_time_series_charts(
        wb_out=wb_out,
        src_sheet=src_sheet,
        ws_center=ws_center,
        n_clusters=n_clusters,
        max_len=max_len,
        labels=labels,
        day_dates=day_dates,
        day_data=day_data,
        used_sheet_names=used_sheet_names,
        chart_count=max_charts,
        y_label="SOC",
        chart_suffix="soc",
        charts_per_row=charts_per_row,
    )


def fill_series(values: List[float], target_len: int) -> List[float]:
    if target_len <= 0:
        return []

    if not values:
        return [0.0] * target_len

    padded = list(values)
    if len(padded) > target_len:
        return padded[:target_len]

    if len(padded) < target_len:
        last = padded[-1]
        padded.extend([last] * (target_len - len(padded)))
    return padded


def load_time_series_by_date(
    wb: openpyxl.Workbook,
    prefer_sheet: str,
    value_col: int,
) -> Dict[str, List[float]]:
    if prefer_sheet and prefer_sheet in wb.sheetnames:
        ws = wb[prefer_sheet]
    else:
        ws = wb[wb.sheetnames[0]]

    tmp: Dict[str, List[Tuple[int, float]]] = defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = parse_date_key(row[0] if len(row) > 0 else None)
        tc = parse_int(row[1] if len(row) > 1 else None)
        v = parse_float(row[value_col] if len(row) > value_col else None)
        if d is None or tc is None or v is None:
            continue
        tmp[d].append((tc, v))

    out: Dict[str, List[float]] = {}
    for d, vals in tmp.items():
        vals = sorted(vals, key=lambda x: x[0])
        out[d] = [x[1] for x in vals]
    return out

def safe_sheet_name(name: str, used: set[str]) -> str:
    base = name[:31]
    if base not in used:
        used.add(base)
        return base
    i = 2
    while True:
        suffix = f"_{i}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        i += 1


def open_with_fallback_copy(path: Path) -> openpyxl.Workbook:
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except PermissionError:
        fallback = Path.home() / f"{path.stem}_copy_for_xmeans{path.suffix}"
        shutil.copy2(path, fallback)
        return openpyxl.load_workbook(fallback, data_only=True)


def xmeans_nd_guarded(
    data: np.ndarray,
    rng: random.Random,
    kmax: int,
    min_bic_gain: float,
    min_child_size: int,
) -> Tuple[np.ndarray, np.ndarray]:
    n, d = data.shape
    if n == 0:
        return np.empty((0,), dtype=int), np.empty((0, d), dtype=float)
    if n == 1:
        return np.zeros(1, dtype=int), data.copy()

    labels, centers = kmeans_nd(data, 1, rng)

    while True:
        k = centers.shape[0]
        if k >= kmax:
            break

        split_any = False
        next_centers: List[np.ndarray] = []

        for c in range(k):
            idx = np.where(labels == c)[0]
            cluster_data = data[idx]

            # Forbid splitting tiny clusters and enforce minimum child size.
            if cluster_data.shape[0] < max(4, min_child_size * 2):
                next_centers.append(centers[c])
                continue

            parent_labels = np.zeros(cluster_data.shape[0], dtype=int)
            parent_centers = np.mean(cluster_data, axis=0, keepdims=True)
            parent_bic = bic_score_nd(cluster_data, parent_labels, parent_centers)

            child_labels, child_centers = kmeans_nd(cluster_data, 2, rng)
            child0 = int(np.sum(child_labels == 0))
            child1 = int(np.sum(child_labels == 1))
            if child0 < min_child_size or child1 < min_child_size:
                next_centers.append(centers[c])
                continue

            child_bic = bic_score_nd(cluster_data, child_labels, child_centers)
            bic_gain = child_bic - parent_bic

            if bic_gain >= min_bic_gain and (len(next_centers) + (k - c) + 1) <= kmax:
                next_centers.extend([child_centers[0], child_centers[1]])
                split_any = True
            else:
                next_centers.append(centers[c])

        if not split_any:
            break

        labels, centers = kmeans_nd(data, len(next_centers), rng)

    return labels, centers


def run(
    input_path: Path,
    output_path: Path,
    seed: int,
    kmax: int,
    min_bic_gain: float,
    min_child_size: int,
    chart_count: int,
    spot_source_path: Optional[Path] = None,
    spot_sheet: Optional[str] = None,
    spot_value_col: int = 14,
    spot_only: bool = False,
) -> Tuple[int, int, Path]:
    wb_in = open_with_fallback_copy(input_path)
    wb_spot = open_with_fallback_copy(spot_source_path) if spot_source_path else None

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)
    used_sheet_names: set[str] = set()

    rng = random.Random(seed)
    total_days = 0
    total_clusters = 0

    for src_sheet in wb_in.worksheets:
        groups: Dict[str, List[Tuple[int, float]]] = {}
        for row in src_sheet.iter_rows(min_row=2, values_only=True):
            dt = parse_date_key(row[0] if len(row) > 0 else None)
            tc = parse_int(row[1] if len(row) > 1 else None)
            val = parse_float(row[2] if len(row) > 2 else None)
            if dt is None or tc is None or val is None:
                continue
            groups.setdefault(dt, []).append((tc, val))

        dates = sorted(groups.keys())
        if not dates:
            continue

        day_vectors: List[List[float]] = []
        day_dates: List[str] = []
        max_len = 0

        for d in dates:
            pts = sorted(groups[d], key=lambda x: x[0])
            vals = [v for _tc, v in pts]
            if not vals:
                continue
            max_len = max(max_len, len(vals))
            day_dates.append(d)
            day_vectors.append(vals)

        if not day_vectors or max_len == 0:
            continue

        for i in range(len(day_vectors)):
            if len(day_vectors[i]) < max_len:
                day_vectors[i] = day_vectors[i] + [day_vectors[i][-1]] * (max_len - len(day_vectors[i]))

        data = np.asarray(day_vectors, dtype=float)
        spot_day_data: Optional[np.ndarray] = None
        if wb_spot is not None:
            prefer_sheet = spot_sheet or src_sheet.title
            spot_series = load_time_series_by_date(wb_spot, prefer_sheet, spot_value_col)
            spot_vectors: List[List[float]] = []
            for d in day_dates:
                spot_vectors.append(fill_series(spot_series.get(d, []), max_len))
            spot_day_data = np.asarray(spot_vectors, dtype=float)

        labels, centers = xmeans_nd_guarded(
            data=data,
            rng=rng,
            kmax=kmax,
            min_bic_gain=min_bic_gain,
            min_child_size=min_child_size,
        )
        n_clusters = int(centers.shape[0])

        # Representative date nearest to each center and member dates per cluster.
        representatives: Dict[int, Tuple[str, float, int]] = {}
        cluster_dates: Dict[int, List[str]] = {}
        for c in range(n_clusters):
            idxs = np.where(labels == c)[0]
            if idxs.size == 0:
                continue
            center = centers[c]
            best_idx = int(idxs[int(np.argmin(np.linalg.norm(data[idxs] - center, axis=1)))])
            representatives[c] = (day_dates[best_idx], float(np.linalg.norm(data[best_idx] - center)), int(idxs.size))
            cluster_dates[c] = [day_dates[i] for i in idxs.tolist()]

        assign_name = safe_sheet_name(f"{src_sheet.title}_assign", used_sheet_names)
        ws_assign = wb_out.create_sheet(assign_name)
        ws_assign.append(
            [
                "sheet",
                "date",
                "cluster_id",
                "distance_to_center",
                "seed",
                "kmax",
                "min_bic_gain",
                "min_child_size",
            ]
        )
        for i, d in enumerate(day_dates):
            c = int(labels[i])
            dist = float(np.linalg.norm(data[i] - centers[c]))
            ws_assign.append([src_sheet.title, d, c, dist, seed, kmax, min_bic_gain, min_child_size])

        repr_name = safe_sheet_name(f"{src_sheet.title}_repr", used_sheet_names)
        ws_repr = wb_out.create_sheet(repr_name)
        ws_repr.append(["sheet", "cluster_id", "representative_date", "member_days", "distance_to_center"])
        for c in range(n_clusters):
            if c in representatives:
                d, dist, members = representatives[c]
                ws_repr.append([src_sheet.title, c, d, members, dist])

        center_name = safe_sheet_name(f"{src_sheet.title}_center", used_sheet_names)
        ws_center = wb_out.create_sheet(center_name)
        ws_center.append(["sheet", "cluster_id"] + [f"t{i+1}" for i in range(max_len)])
        for c in range(n_clusters):
            ws_center.append([src_sheet.title, c] + [float(x) for x in centers[c].tolist()])

        dates_name = safe_sheet_name(f"{src_sheet.title}_dates", used_sheet_names)
        ws_dates = wb_out.create_sheet(dates_name)
        ws_dates.append(["sheet", "cluster_id", "member_index", "date"])
        for c in range(n_clusters):
            for i, d in enumerate(cluster_dates.get(c, []), start=1):
                ws_dates.append([src_sheet.title, c, i, d])

        if not spot_only:
            add_soc_charts(
                wb_out=wb_out,
                src_sheet=src_sheet,
                ws_center=ws_center,
                n_clusters=n_clusters,
                max_len=max_len,
                labels=labels,
                day_dates=day_dates,
                day_data=data,
                used_sheet_names=used_sheet_names,
                max_charts=chart_count,
                charts_per_row=3,
            )

        if spot_day_data is not None:
            add_cluster_time_series_charts(
                wb_out=wb_out,
                src_sheet=src_sheet,
                ws_center=ws_center,
                n_clusters=n_clusters,
                max_len=max_len,
                labels=labels,
                day_dates=day_dates,
                day_data=spot_day_data,
                used_sheet_names=used_sheet_names,
                chart_count=chart_count,
                y_label="Spot",
                chart_suffix="spot",
                charts_per_row=3,
            )

        total_days += len(day_dates)
        total_clusters += n_clusters

    actual_output_path = output_path
    try:
        wb_out.save(actual_output_path)
    except PermissionError as exc:
        actual_output_path = output_path.with_name(f"{output_path.stem}_retry_{datetime.now().strftime('%Y%m%d_%H%M%S')}{output_path.suffix}")
        wb_out.save(actual_output_path)
        print(f"PermissionError: {exc}")
        print(f"Original output is in use. Saved to: {actual_output_path}")

    wb_in.close()
    if wb_spot is not None:
        wb_spot.close()

    return total_days, total_clusters, actual_output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="X-means clustering for daily value transition profiles (A=date, B=timecode, C=value)."
    )
    parser.add_argument("--input", dest="input_path", default=None, help="Input xlsx path")
    parser.add_argument("--output", dest="output_path", default=None, help="Output xlsx path")
    parser.add_argument("--seed", type=int, default=42, help="Random seed (default: 42)")
    parser.add_argument("--kmax", type=int, default=50, help="Maximum clusters (default: 50)")
    parser.add_argument(
        "--min-bic-gain",
        type=float,
        default=200.0,
        help="Minimum BIC improvement required to accept a split (default: 200.0)",
    )
    parser.add_argument(
        "--min-child-size",
        type=int,
        default=1,
        help="Minimum number of days in each child cluster after split (default: 1)",
    )
    parser.add_argument(
        "--chart-count",
        type=int,
        default=13,
        help="Maximum number of SOC charts to generate per source sheet (default: 13)",
    )
    parser.add_argument("--spot-source", dest="spot_source", default=None, help="Raw workbook used for spot price series")
    parser.add_argument("--spot-sheet", dest="spot_sheet", default=None, help="Sheet name in spot source workbook")
    parser.add_argument(
        "--spot-value-col",
        type=int,
        default=14,
        help="0-based column index for spot price value in --spot-source (default: 14)",
    )
    parser.add_argument(
        "--spot-only",
        action="store_true",
        help="Generate only spot charts (skip SOC charts)",
    )
    args = parser.parse_args()

    if args.spot_only and not args.spot_source:
        raise SystemExit("--spot-only requires --spot-source to be specified.")

    if args.input_path:
        input_path = Path(args.input_path)
    else:
        found = find_file_under_onedrive(DEFAULT_INPUT_NAME)
        if found is None:
            raise FileNotFoundError(f"Input file not found: {DEFAULT_INPUT_NAME}")
        input_path = found

    if args.output_path:
        output_path = Path(args.output_path)
    else:
        parent = input_path.parent
        output_path = parent / f"{input_path.stem}_xmeans_daytrend_seed{args.seed}_kmax{args.kmax}.xlsx"

    spot_source = Path(args.spot_source) if args.spot_source else None

    days, clusters, actual_output_path = run(
        input_path=input_path,
        output_path=output_path,
        seed=args.seed,
        kmax=args.kmax,
        min_bic_gain=args.min_bic_gain,
        min_child_size=args.min_child_size,
        chart_count=args.chart_count,
        spot_source_path=spot_source,
        spot_sheet=args.spot_sheet,
        spot_value_col=args.spot_value_col,
        spot_only=args.spot_only,
    )
    print(f"input={input_path}")
    print(f"output={actual_output_path}")
    print(f"seed={args.seed}")
    print(f"kmax={args.kmax}")
    print(f"min_bic_gain={args.min_bic_gain}")
    print(f"min_child_size={args.min_child_size}")
    print(f"days={days}, total_clusters={clusters}")


if __name__ == "__main__":
    main()

import argparse
import math
import os
import random
import re
from typing import List, Optional, Sequence, Tuple

import numpy as np
import openpyxl


DEFAULT_INPUT_NAME = "extracted_A_L_yyyymm_split_1-0_2-1_D_normalized_minmax.xlsx"


def find_file_under_onedrive(filename: str) -> Optional[str]:
    start = r"C:\Users\kit02\OneDrive"
    for root, _, files in os.walk(start):
        if filename in files:
            return os.path.join(root, filename)
    return None


def parse_date(c_value: object) -> Optional[str]:
    if c_value is None:
        return None
    text = str(c_value).strip()
    m = re.match(r"^(\d{8})", text)
    return m.group(1) if m else None


def parse_float(v: object) -> Optional[float]:
    if v is None:
        return None
    text = str(v).strip()
    if text == "":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def kmeans_1d(data: Sequence[float], k: int, rng: random.Random, max_iter: int = 100) -> Tuple[List[int], List[float]]:
    n = len(data)
    if n == 0:
        return [], []
    if k <= 1:
        return [0] * n, [sum(data) / n]

    sorted_vals = sorted(data)
    centers = [sorted_vals[int((i + 0.5) * n / k)] for i in range(k)]

    labels = [0] * n
    for _ in range(max_iter):
        changed = False

        for i, x in enumerate(data):
            best = min(range(k), key=lambda c: abs(x - centers[c]))
            if labels[i] != best:
                labels[i] = best
                changed = True

        new_centers: List[float] = []
        for c in range(k):
            pts = [data[i] for i in range(n) if labels[i] == c]
            if pts:
                new_centers.append(sum(pts) / len(pts))
            else:
                new_centers.append(rng.choice(list(data)))

        if all(abs(new_centers[i] - centers[i]) < 1e-12 for i in range(k)) and not changed:
            centers = new_centers
            break
        centers = new_centers

    # Stable output: reorder clusters by center value.
    order = sorted(range(k), key=lambda i: centers[i])
    remap = {old: new for new, old in enumerate(order)}
    centers = [centers[i] for i in order]
    labels = [remap[l] for l in labels]

    return labels, centers


def sse_1d(data: Sequence[float], labels: Sequence[int], centers: Sequence[float]) -> float:
    return sum((x - centers[labels[i]]) ** 2 for i, x in enumerate(data))


def bic_score(data: Sequence[float], labels: Sequence[int], centers: Sequence[float]) -> float:
    n = len(data)
    k = len(centers)
    if n == 0 or k == 0:
        return float("-inf")

    sse = sse_1d(data, labels, centers)
    dof = max(n - k, 1)
    var = max(sse / dof, 1e-12)

    counts = [0] * k
    for label in labels:
        counts[label] += 1

    c0 = -0.5 * math.log(2.0 * math.pi * var)
    log_likelihood = 0.0
    for i, x in enumerate(data):
        label = labels[i]
        pi_k = max(counts[label] / n, 1e-12)
        diff = x - centers[label]
        log_likelihood += c0 - (diff * diff) / (2.0 * var) + math.log(pi_k)

    # params: k means + (k - 1) weights + 1 shared variance
    p = k + (k - 1) + 1
    return log_likelihood - 0.5 * p * math.log(n)


def xmeans_1d(data: Sequence[float], rng: random.Random, kmax: int = 8) -> Tuple[List[int], List[float]]:
    n = len(data)
    if n == 0:
        return [], []
    if n == 1:
        return [0], [data[0]]

    labels, centers = kmeans_1d(data, 1, rng)

    while True:
        k = len(centers)
        if k >= kmax:
            break

        split_any = False
        next_centers: List[float] = []

        for c in range(k):
            idxs = [i for i, label in enumerate(labels) if label == c]
            cluster_data = [data[i] for i in idxs]

            if len(cluster_data) < 4:
                next_centers.append(centers[c])
                continue

            parent_labels = [0] * len(cluster_data)
            parent_centers = [sum(cluster_data) / len(cluster_data)]
            parent_bic = bic_score(cluster_data, parent_labels, parent_centers)

            child_labels, child_centers = kmeans_1d(cluster_data, 2, rng)
            child_bic = bic_score(cluster_data, child_labels, child_centers)

            if child_bic > parent_bic and (len(next_centers) + (k - c) + 1) <= kmax:
                next_centers.extend(child_centers)
                split_any = True
            else:
                next_centers.append(centers[c])

        if not split_any:
            break

        labels, centers = kmeans_1d(data, len(next_centers), rng)

    return labels, centers


def nearest_points_to_centroids_1d(
    data: Sequence[float], labels: Sequence[int], centers: Sequence[float]
) -> List[Tuple[int, int, float, float]]:
    """Return one nearest real data point for each cluster centroid."""
    if not data or not centers:
        return []

    arr = np.asarray(data, dtype=float)
    lbl = np.asarray(labels, dtype=int)
    ctr = np.asarray(centers, dtype=float)

    representatives: List[Tuple[int, int, float, float]] = []
    for cluster_id, center in enumerate(ctr):
        cluster_indices = np.where(lbl == cluster_id)[0]
        if cluster_indices.size == 0:
            continue
        distances = np.abs(arr[cluster_indices] - center)
        nearest_pos = int(cluster_indices[int(np.argmin(distances))])
        representatives.append(
            (cluster_id, nearest_pos, float(center), float(arr[nearest_pos]))
        )
    return representatives


def unique_sheet_name(workbook: openpyxl.Workbook, desired: str) -> str:
    name = desired[:31]
    if name not in workbook.sheetnames:
        return name

    i = 1
    while True:
        suffix = f"_{i}"
        cand = desired[: 31 - len(suffix)] + suffix
        if cand not in workbook.sheetnames:
            return cand
        i += 1


def run(input_path: str, output_path: str, summary_path: str, seed: int, kmax: int) -> Tuple[int, int]:
    rng = random.Random(seed)

    src = openpyxl.load_workbook(input_path, data_only=True)
    out = openpyxl.Workbook()
    out.remove(out[out.sheetnames[0]])

    summary_lines = ["sheet,date,rows_total,rows_numeric_d,clusters,seed,kmax"]
    representative_lines = [
        "sheet,date,cluster_id,centroid_value,nearest_d_value,row_index_in_date_group,row_number_in_output_sheet"
    ]
    created_sheets = 0

    for source_sheet_name in src.sheetnames:
        ws = src[source_sheet_name]
        groups = {}

        for row in ws.iter_rows(values_only=True):
            date = parse_date(row[2] if len(row) >= 3 else None)
            if date is None:
                continue
            groups.setdefault(date, []).append(list(row))

        for date in sorted(groups.keys()):
            rows = groups[date]

            numeric_row_indexes: List[int] = []
            values: List[float] = []

            for idx, row in enumerate(rows):
                d_val = parse_float(row[3] if len(row) >= 4 else None)
                if d_val is None:
                    continue
                numeric_row_indexes.append(idx)
                values.append(d_val)

            if values:
                labels, centers = xmeans_1d(values, rng=rng, kmax=kmax)
                n_clusters = len(centers)
                representatives = nearest_points_to_centroids_1d(values, labels, centers)
            else:
                labels = []
                centers = []
                n_clusters = 0
                representatives = []

            row_cluster = [None] * len(rows)
            for i, row_index in enumerate(numeric_row_indexes):
                row_cluster[row_index] = int(labels[i])

            out_sheet_name = unique_sheet_name(out, f"{source_sheet_name}_{date}")
            ows = out.create_sheet(out_sheet_name)
            for i, row in enumerate(rows):
                ows.append(row + [row_cluster[i]])
            created_sheets += 1

            summary_lines.append(
                f"{source_sheet_name},{date},{len(rows)},{len(values)},{n_clusters},{seed},{kmax}"
            )
            for cluster_id, numeric_pos, centroid_value, nearest_value in representatives:
                row_index_in_date_group = numeric_row_indexes[numeric_pos]
                row_number_in_output_sheet = row_index_in_date_group + 1
                representative_lines.append(
                    f"{source_sheet_name},{date},{cluster_id},{centroid_value},{nearest_value},"
                    f"{row_index_in_date_group},{row_number_in_output_sheet}"
                )

    out.save(output_path)
    with open(summary_path, "w", encoding="utf-8", newline="") as f:
        f.write("\n".join(summary_lines) + "\n")
    rep_path = summary_path.replace("_summary.csv", "_cluster_representatives.csv")
    if rep_path == summary_path:
        rep_path = summary_path + ".cluster_representatives.csv"
    with open(rep_path, "w", encoding="utf-8", newline="") as f:
        f.write("\n".join(representative_lines) + "\n")

    return created_sheets, len(summary_lines) - 1


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Split rows by date from column C (first 8 digits: yyyymmdd) and "
            "run deterministic X-means clustering on column D."
        )
    )
    parser.add_argument("--input", dest="input_path", default=None, help="Input xlsx path")
    parser.add_argument("--output", dest="output_path", default=None, help="Output xlsx path")
    parser.add_argument("--summary", dest="summary_path", default=None, help="Summary csv path")
    parser.add_argument("--seed", type=int, default=42, help="Random seed for reproducibility")
    parser.add_argument("--kmax", type=int, default=8, help="Maximum number of clusters")
    args = parser.parse_args()

    input_path = args.input_path
    if input_path is None:
        found = find_file_under_onedrive(DEFAULT_INPUT_NAME)
        if found is None:
            raise SystemExit(f"Input not found under OneDrive: {DEFAULT_INPUT_NAME}")
        input_path = found

    base = os.path.dirname(input_path)
    output_path = args.output_path or os.path.join(
        base, f"extracted_A_L_yyyymm_split_1-0_2-1_D_norm_by_date_xmeans_seed{args.seed}.xlsx"
    )
    summary_path = args.summary_path or os.path.join(
        base, f"extracted_A_L_yyyymm_split_1-0_2-1_D_norm_by_date_xmeans_seed{args.seed}_summary.csv"
    )

    created_sheets, group_count = run(
        input_path=input_path,
        output_path=output_path,
        summary_path=summary_path,
        seed=args.seed,
        kmax=args.kmax,
    )

    print(f"input={input_path}")
    print(f"output={output_path}")
    print(f"summary={summary_path}")
    print(f"seed={args.seed}")
    print(f"kmax={args.kmax}")
    print(f"groups={group_count}")
    print(f"created_sheets={created_sheets}")


if __name__ == "__main__":
    main()

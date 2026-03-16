from __future__ import annotations

from pathlib import Path
import openpyxl

INPUT_PATH = Path(r"C:\Users\kit02\spot_summary_2024_copy.xlsx")
OUTPUT_PATH = Path(r"C:\Users\kit02\OneDrive\研究\エクセル\x-means_\spot_summary_2024_ABO_O_normalized_mean48.xlsx")
BLOCK_SIZE = 48


def to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def normalize_block_by_mean(values):
    nums = [x for x in values if x is not None]
    if not nums:
        return [None for _ in values]
    mean = sum(nums) / len(nums)
    if mean == 0:
        return [0.0 if x is not None else None for x in values]
    return [(x / mean) if x is not None else None for x in values]


wb_in = openpyxl.load_workbook(INPUT_PATH, data_only=True)
wb_out = openpyxl.Workbook()
wb_out.remove(wb_out.active)

for sheet_name in wb_in.sheetnames:
    ws_in = wb_in[sheet_name]
    ws_out = wb_out.create_sheet(title=sheet_name[:31])

    a_header = ws_in.cell(row=1, column=1).value
    b_header = ws_in.cell(row=1, column=2).value
    o_header = ws_in.cell(row=1, column=15).value
    ws_out.cell(row=1, column=1, value=a_header if a_header is not None else "A")
    ws_out.cell(row=1, column=2, value=b_header if b_header is not None else "B")
    ws_out.cell(
        row=1,
        column=3,
        value=(f"{o_header}_normalized_mean48" if o_header is not None else "O_normalized_mean48"),
    )

    records = []
    for r in range(2, ws_in.max_row + 1):
        a = ws_in.cell(row=r, column=1).value
        b = ws_in.cell(row=r, column=2).value
        o_raw = ws_in.cell(row=r, column=15).value
        records.append((a, b, to_float(o_raw)))

    normalized_o = []
    for i in range(0, len(records), BLOCK_SIZE):
        block = records[i : i + BLOCK_SIZE]
        block_o = [x[2] for x in block]
        normalized_o.extend(normalize_block_by_mean(block_o))

    for i, (a, b, _o) in enumerate(records, start=2):
        ws_out.cell(row=i, column=1, value=a)
        ws_out.cell(row=i, column=2, value=b)
        ws_out.cell(row=i, column=3, value=normalized_o[i - 2])

wb_out.save(OUTPUT_PATH)
print(f"written: {OUTPUT_PATH}")

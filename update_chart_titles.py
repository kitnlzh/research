import pandas as pd
import openpyxl
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font

# Load the Excel file
file_path = r"C:\Users\kit02\OneDrive\研究\エクセル\x-means_\spot_summary_2024_ABO_O_normalized_mean48_xmeans_daytrend_seed42_kmax50_spot13_yaxis_fix_20260317.xlsx"

# Load the workbook
wb = openpyxl.load_workbook(file_path)

# Read the dates sheet to get cluster-date mapping
dates_sheet = "spot_summary_2024 (1)_dates"
df_dates = pd.read_excel(file_path, sheet_name=dates_sheet)

# Create cluster_id to date mapping
cluster_date_map = {}
for _, row in df_dates.iterrows():
    cluster_date_map[int(row['cluster_id'])] = str(row['date'])

print("Cluster to Date mapping:")
for cluster_id, date in cluster_date_map.items():
    print(f"Cluster {cluster_id}: {date}")

# Check all sheets for charts
print("\nChecking all sheets for charts:")
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    if hasattr(ws, '_charts') and ws._charts:
        print(f"Sheet '{sheet_name}' has {len(ws._charts)} chart(s):")
        for i, chart in enumerate(ws._charts):
            print(f"  Chart {i}: title='{chart.title}'")
    else:
        print(f"Sheet '{sheet_name}': no charts found")

# Read the spot sheet
spot_sheet = "spot_summary_2024 (1)_spot"
df_spot = pd.read_excel(file_path, sheet_name=spot_sheet)

print(f"\nSpot sheet data:")
print(df_spot)

# If charts are found, update them
charts_updated = 0
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    if hasattr(ws, '_charts') and ws._charts:
        for chart in ws._charts:
            if hasattr(chart, 'title') and chart.title:
                title_text = str(chart.title)
                print(f"Processing chart with title: '{title_text}'")
                # Try to extract cluster id from title
                if "plot" in title_text:
                    try:
                        cluster_id = int(title_text.split()[-1])
                        if cluster_id in cluster_date_map:
                            new_title = f"Cluster {cluster_id}: {cluster_date_map[cluster_id]}"
                            chart.title = new_title
                            print(f"Updated chart title to: {new_title}")
                            charts_updated += 1
                    except ValueError:
                        print(f"Could not extract cluster ID from title: '{title_text}'")
                elif title_text.isdigit():
                    try:
                        cluster_id = int(title_text)
                        if cluster_id in cluster_date_map:
                            new_title = f"Cluster {cluster_id}: {cluster_date_map[cluster_id]}"
                            chart.title = new_title
                            print(f"Updated chart title to: {new_title}")
                            charts_updated += 1
                    except ValueError:
                        print(f"Could not parse cluster ID from title: '{title_text}'")

print(f"\nTotal charts updated: {charts_updated}")

# Save the workbook
wb.save(file_path)
print(f"\nWorkbook saved to {file_path}")